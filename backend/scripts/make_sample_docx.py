"""Create Word (.docx) versions of the Markdown samples in samples/.

Run from the backend folder:  .\.venv\Scripts\python.exe scripts\make_sample_docx.py
Handles the simple Markdown used in the samples: # headings, - bullets, | tables |, paragraphs.
"""

import sys
from pathlib import Path

import docx

SAMPLES = Path(__file__).resolve().parents[2] / "samples"


def markdown_to_docx(md_path: Path, docx_path: Path) -> None:
    document = docx.Document()
    table_rows: list[list[str]] = []

    def flush_table():
        rows = [r for r in table_rows if not all(set(c) <= set("-: ") for c in r)]  # drop |---| lines
        if rows:
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        table_rows.clear()

    for line in md_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("|"):
            table_rows.append([cell.strip() for cell in line.strip().strip("|").split("|")])
            continue
        flush_table()
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("#").strip()
            if level == 1:
                document.add_heading(text, level=0)  # Word's "Title" style
            else:
                document.add_heading(text, level=level - 1)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            document.add_paragraph(line)
    flush_table()
    document.save(docx_path)


if __name__ == "__main__":
    for md in sorted(SAMPLES.glob("*.md")):
        out = md.with_suffix(".docx")
        markdown_to_docx(md, out)
        print(f"Created {out.name}")
    sys.exit(0)
