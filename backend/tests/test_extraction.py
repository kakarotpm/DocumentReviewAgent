import io

import docx
import openpyxl
import pytest

from app.config import PROJECT_ROOT
from app.services.extraction import MAX_ROWS_PER_SHEET, ExtractionError, extract_text, extract_workbook
from tests.pdf_helper import make_pdf

SAMPLES = PROJECT_ROOT / "samples"
LONG_LINE = "The scope lists in-scope items and out-of-scope items for the project."


def make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("My BRD", level=0)
    document.add_heading("3. Scope", level=1)
    document.add_paragraph(LONG_LINE)
    document.add_paragraph("Supplier invoices", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "ID", "Requirement"
    table.cell(1, 0).text, table.cell(1, 1).text = "BR-01", "Match invoices to POs"
    document.add_heading("3.1 Out of scope", level=2)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------- Each format ----------


def test_docx_keeps_headings_lists_and_tables_in_order():
    file_type, text = extract_text("brd.docx", make_docx())
    assert file_type == "docx"
    assert "# My BRD" in text
    assert "# 3. Scope" in text
    assert "## 3.1 Out of scope" in text
    assert "- Supplier invoices" in text
    assert "| BR-01 | Match invoices to POs |" in text
    assert text.index("3. Scope") < text.index("BR-01") < text.index("3.1 Out of scope")


def test_pdf_text_with_page_markers():
    file_type, text = extract_text("brd.pdf", make_pdf(["Business Requirements Document", LONG_LINE]))
    assert file_type == "pdf"
    assert text.startswith("[Page 1]")
    assert LONG_LINE in text


def test_markdown_and_text_files():
    assert extract_text("a.md", f"# Title\n\n{LONG_LINE}".encode())[1].startswith("# Title")
    assert extract_text("a.TXT", LONG_LINE.encode())[0] == "txt"  # extension case doesn't matter


def test_text_file_with_windows_encoding():
    _, text = extract_text("a.txt", f"Budget: £180,000. {LONG_LINE}".encode("cp1252"))
    assert "£180,000" in text


@pytest.mark.parametrize("name", ["brd_strong.md", "brd_weak.md", "brd_strong.docx", "brd_weak.docx"])
def test_sample_files_extract(name):
    _, text = extract_text(name, (SAMPLES / name).read_bytes())
    assert "Scope" in text


# ---------- Friendly errors ----------


def error_for(filename, data):
    with pytest.raises(ExtractionError) as info:
        extract_text(filename, data)
    return str(info.value)


def test_unsupported_file_type():
    assert "Please upload a .docx, .pdf, .md or .txt file" in error_for("diagram.png", b"x" * 100)


def test_old_doc_format_is_unsupported():
    assert "not a supported file type" in error_for("old.doc", b"x" * 100)


def test_empty_file():
    assert "is empty" in error_for("a.txt", b"")


def test_file_too_large():
    assert "The limit is 10 MB" in error_for("a.txt", b"x" * (10 * 1024 * 1024 + 1))


def test_scanned_pdf():
    message = error_for("scan.pdf", make_pdf([]))
    assert "looks like a scanned PDF" in message
    assert "upload the original Word file" in message


def test_damaged_docx():
    assert "couldn't be opened as a Word document" in error_for("bad.docx", b"not really a word file" * 10)


def test_damaged_pdf():
    assert "couldn't be read as a PDF" in error_for("bad.pdf", b"not really a pdf" * 10)


def test_nearly_empty_text_file():
    assert "almost no text" in error_for("a.txt", b"hello")


# ---------- Excel workbooks ----------


def make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_workbook_sheets_headers_and_rows():
    data = make_xlsx({
        "Field Mapping": [["Target Field", "Source Field", "Transformation"], ["invoice_id", "BELNR", None], [], ["amount", "WRBTR", "  round to   2dp "]],
        "Notes": [],
    })
    text = extract_workbook("mapping.xlsx", data)
    assert text.startswith("Workbook: mapping.xlsx (2 sheets)")
    assert "## Sheet: Field Mapping\n(3 rows with content)\n| Target Field | Source Field | Transformation |" in text
    assert "| invoice_id | BELNR |\n" in text  # trailing empty cells dropped, empty rows skipped
    assert "| amount | WRBTR | round to 2dp |" in text  # whitespace tidied
    assert "## Sheet: Notes\n(empty sheet)" in text


def test_workbook_keeps_only_the_first_rows_of_a_long_sheet():
    data = make_xlsx({"Big": [["Row", n] for n in range(MAX_ROWS_PER_SHEET + 5)]})
    text = extract_workbook("big.xlsx", data)
    assert f"{MAX_ROWS_PER_SHEET + 5} rows with content; only the first {MAX_ROWS_PER_SHEET} are shown" in text
    assert f"| Row | {MAX_ROWS_PER_SHEET - 1} |" in text
    assert f"| Row | {MAX_ROWS_PER_SHEET} |" not in text


def test_sample_workbook_extracts():
    text = extract_workbook("frd_mapping_sample.xlsx", (SAMPLES / "frd_mapping_sample.xlsx").read_bytes())
    for sheet in ["Read Me", "Source Systems", "Field Mapping", "Joins", "Measures", "Reconciliation"]:
        assert f"## Sheet: {sheet}" in text


@pytest.mark.parametrize(
    "filename, data, message",
    [
        ("old.xls", b"x" * 100, "save it as .xlsx"),
        ("mapping.csv", b"a,b", "Please upload a .xlsx file"),
        ("empty.xlsx", b"", "is empty"),
        ("broken.xlsx", b"not a zip file", "couldn't be opened as an Excel workbook"),
        ("blank.xlsx", make_xlsx({"Sheet1": []}), "has no content in any sheet"),
    ],
)
def test_unusable_workbooks(filename, data, message):
    with pytest.raises(ExtractionError, match=message):
        extract_workbook(filename, data)
